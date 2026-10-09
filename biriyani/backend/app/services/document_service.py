from __future__ import annotations
import io
import pypdf
import docx
from sqlalchemy.orm import Session

from app.models.conversation import Conversation
from app.models.document import Document

import re

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB limit


class DocumentProcessingError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message


def normalize_text(text: str) -> str:
    """Normalize extracted text by cleaning up control characters and excessive whitespace."""
    if not text:
        return ""
    # Strip invalid control characters
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    # Replace carriage returns
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    # Replace multiple consecutive blank lines with a double newline
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    # Remove excessive trailing whitespace per line
    lines = [line.rstrip() for line in cleaned.split("\n")]
    return "\n".join(lines).strip()


def extract_text_from_file(filename: str, content: bytes) -> tuple[str, str]:
    if len(content) > MAX_FILE_SIZE:
        raise DocumentProcessingError(
            "FILE_TOO_LARGE", f"File size exceeds maximum allowed limit of {MAX_FILE_SIZE // (1024 * 1024)}MB."
        )

    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext == "pdf":
        file_type = "pdf"
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            extracted_pages = []
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    extracted_pages.append(txt)
            text = "\n".join(extracted_pages)
        except Exception as e:
            raise DocumentProcessingError(
                "PDF_EXTRACTION_FAILED", f"Could not extract text from PDF document: {str(e)}"
            ) from e

    elif ext in ("docx", "doc"):
        file_type = "docx"
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
        except Exception as e:
            raise DocumentProcessingError(
                "DOCX_EXTRACTION_FAILED", f"Could not extract text from DOCX document: {str(e)}"
            ) from e

    elif ext in ("csv", "tsv"):
        file_type = "csv"
        try:
            raw_str = content.decode("utf-8", errors="replace")
            text = raw_str
        except Exception as e:
            raise DocumentProcessingError(
                "CSV_EXTRACTION_FAILED", f"Could not extract text from CSV file: {str(e)}"
            ) from e

    elif ext in ("txt", "text", "md", "json", "log"):
        file_type = "txt"
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="ignore")

    else:
        raise DocumentProcessingError(
            "UNSUPPORTED_FILE_TYPE", "Only PDF, DOCX, TXT, MD, CSV, and JSON files are supported."
        )

    normalized = normalize_text(text)
    if not normalized:
        raise DocumentProcessingError(
            "EMPTY_DOCUMENT_TEXT", "The uploaded file contained no extractable text or content."
        )

    return file_type, normalized



def save_document(
    db: Session,
    conversation: Conversation,
    filename: str,
    file_type: str,
    file_size: int,
    extracted_text: str,
) -> Document:
    doc = Document(
        conversation_id=conversation.id,
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        extracted_text=extracted_text,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    if conversation.user_id:
        from app.services.event_service import ApplicationEventService
        ApplicationEventService.record_event(
            db,
            event_type="DOCUMENT_UPLOADED",
            user_id=conversation.user_id,
            entity_type="document",
            entity_id=doc.id,
            metadata={"filename": filename, "file_type": file_type, "file_size": file_size},
            event_id=f"doc-upload-{doc.id}",
        )
        ApplicationEventService.record_event(
            db,
            event_type="DOCUMENT_PROCESSING_COMPLETED",
            user_id=conversation.user_id,
            entity_type="document",
            entity_id=doc.id,
            metadata={"filename": filename, "char_count": len(extracted_text)},
            event_id=f"doc-proc-{doc.id}",
        )
        db.commit()

    return doc


def get_conversation_documents(db: Session, conversation_id: str) -> list[Document]:
    return (
        db.query(Document)
        .filter(Document.conversation_id == conversation_id)
        .order_by(Document.created_at.asc())
        .all()
    )


def get_document(db: Session, document_id: str) -> Document | None:
    return db.get(Document, document_id)


def delete_document(db: Session, document: Document) -> None:
    user_id = document.conversation.user_id if document.conversation else None
    doc_id = document.id
    filename = document.filename

    db.delete(document)
    db.commit()

    if user_id:
        from app.services.event_service import ApplicationEventService
        ApplicationEventService.record_event(
            db,
            event_type="DOCUMENT_DELETED",
            user_id=user_id,
            entity_type="document",
            entity_id=doc_id,
            metadata={"filename": filename},
            event_id=f"doc-del-{doc_id}",
        )
        db.commit()

